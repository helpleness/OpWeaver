"""Export evidence-bound experiment report and searchable paired-answer appendix.

Run with the bundled document Python runtime. No model execution or new scoring.
"""

import argparse
import html
import json
import shutil
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
RUNS = ROOT / "runs"
MODES = {
    "base_full_history": "原始模型 完整历史",
    "base_recent_window": "原始模型 最近窗口",
    "base_same_event_text": "原始模型 同事件原文",
    "base_event_text": "原始模型 同段原文",
    "old_memory": "旧记忆版",
    "episode_small": "新记忆 小数据版",
    "episode_open": "新记忆 公开语料版",
    "episode_open_grouped": "新记忆 公开语料及归组版",
    "episode_memory": "新记忆 回归测试",
}
LABELS = dict(
    correct="目标正确",
    partial="部分对应",
    not_answered="目标未答",
    unsupported_answer="目标无依据或错误",
    unresolved_conflict="擅自解决未决冲突",
    uncertain="无法判定",
    acceptable="整句合格",
    problematic="有问题",
    honest_but_incomplete="诚实但不完整",
    unjudgeable="无法判定",
)


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def rows(path):
    return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines()]


def joined(suite, review_dir, source_dir):
    mapping = read(review_dir / "private_mapping.json")
    packets = {p["review_id"]: p for p in rows(review_dir / "all_packets.jsonl")}
    outputs = {}
    if suite == "long":
        for seed in (57511, 57512):
            for r in rows(source_dir / f"seed_{seed}" / "outputs.jsonl"):
                outputs[seed, r["id"], r["mode"]] = r
    else:
        outputs = {(r["id"], r["mode"]): r for r in rows(source_dir / "outputs.jsonl")}
    result = []
    for r in rows(review_dir / "final_reviews.jsonl"):
        p, m = packets[r["review_id"]], mapping[r["review_id"]]
        assert p["binding_sha256"] == r["binding_sha256"]
        oid = m.get("output_id", m.get("case_id"))
        key = (m["seed"], oid, m["mode"]) if suite == "long" else (oid, m["mode"])
        output = outputs[key]
        if output["result"].get("answer") is not None:
            assert output["result"]["answer"] == p["answer"]
        result.append(
            dict(
                suite=suite,
                **m,
                case_key=f"{suite}:{m.get('seed', '')}:{oid}",
                review_id=r["review_id"],
                packet=p,
                review=r,
                original_output=output,
            )
        )
    return result


def summary(records):
    result = {}
    for mode in sorted({r["mode"] for r in records}):
        rs = [r["review"] for r in records if r["mode"] == mode]
        result[mode] = dict(
            n=len(rs),
            target=dict(Counter(r["target"] for r in rs)),
            whole=dict(Counter(r["whole_reply"] for r in rs)),
            event_errors=sum(r["event_claim_problem"] for r in rs),
            status_errors=sum(r["status_problem"] for r in rs),
        )
    return result


def paired(records, baseline, modified="episode_open_grouped"):
    lookup = {(r["case_key"], r["mode"]): r for r in records}
    count = Counter()
    for key in sorted({r["case_key"] for r in records}):
        a, b = lookup[key, baseline], lookup[key, modified]
        good_a = a["review"]["whole_reply"] == "acceptable"
        good_b = b["review"]["whole_reply"] == "acceptable"
        count[
            "both_good"
            if good_a and good_b
            else "new_only"
            if good_b
            else "base_only"
            if good_a
            else "both_bad"
        ] += 1
    return {k: count[k] for k in ("both_good", "new_only", "base_only", "both_bad")}


def appendix(records, path):
    e = html.escape
    suites = dict(
        long="主实验 长上下文",
        fresh="辅助 新事件七组",
        public="辅助 真实对话",
        regression="旧问题回归",
    )
    groups = defaultdict(list)
    for r in records:
        groups[r["case_key"]].append(r)
    parts = []
    for key, group in groups.items():
        first = group[0]
        title = (
            f"{suites[first['suite']]} | {key} | {first.get('budget', first.get('checkpoint', ''))}"
        )
        body = [
            f'<section class="case" data-suite="{e(first["suite"])}"><h2>{e(title)}</h2><p class="question">{e(first["packet"]["question"])}</p>'
        ]
        for r in sorted(group, key=lambda x: x["mode"]):
            p, j, result = r["packet"], r["review"], r["original_output"]["result"]
            tag = (
                "good"
                if j["whole_reply"] == "acceptable"
                else "bad"
                if j["whole_reply"] == "problematic"
                else "partial"
            )
            body.append(
                f'<article class="answer {tag}"><h3>{e(MODES[r["mode"]])}</h3><p class="label">{e(LABELS[j["target"]])} · {e(LABELS[j["whole_reply"]])} · ID {e(r["review_id"])}</p><blockquote>{e(p["answer"])}</blockquote>'
            )
            body.append(
                f"<p><b>复核理由</b> {e(j['rationale'])}</p><p>事实问题 {j['event_claim_problem']} · 状态问题 {j['status_problem']} · 输入 token {result.get('input_tokens', '未记录')}</p>"
            )
            body.append(
                f"<details><summary>查看证据和实际读取范围</summary><p>{e(p['execution_state'])}</p>"
            )
            if not p["evidence"]:
                body.append("<p>该条件无相关事件证据。</p>")
            for evidence in p["evidence"]:
                body.append(f"<p><b>{e(evidence['id'])}</b> {e(evidence['text'])}</p>")
            body.append(
                f'<p class="hash">绑定 SHA256 {e(p["binding_sha256"])}</p></details></article>'
            )
        body.append("</section>")
        parts.append("\n".join(body))
    page = (
        """<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>逐题对照与证据附录</title>
<style>body{font-family:'Microsoft YaHei',sans-serif;background:#f4f6f8;color:#172330;max-width:1100px;margin:auto;padding:28px;line-height:1.7}h1,h2,h3{line-height:1.4}header{position:sticky;top:0;background:#fff;padding:15px 20px;border:1px solid #ccd3da;z-index:1}input,select{font:inherit;padding:8px;margin:4px;border:1px solid #aab6c2}.case{background:white;margin:24px 0;padding:24px;border:1px solid #ccd3da}h2{font-size:20px}.question{font-weight:bold}.answer{border-top:1px solid #d6dde5;padding:12px 0}.label{font-size:14px;color:#506070}.good .label{color:#176543}.bad .label{color:#9c3030}blockquote{margin:8px 0;padding:12px;background:#f5f7fa;white-space:pre-wrap}details{font-size:14px}.hash{word-break:break-all;color:#687482}h3{font-size:17px}#count{color:#506070}@media print{header{position:static}details{display:block}.case{break-before:page}}</style>
<h1>逐题对照与证据附录</h1><p>同一 case 内的问题与历史检查点配对。不同系统的实际输入与读取范围可能不同，证据以各回答下方的记录为准。标签为 AI 初审及主助手证据复核，非人工金标准。没有按固定拒答语判分。</p>
<header><input id="search" placeholder="搜索问题、回答、证据或 ID"><select id="suite"><option value="">全部实验</option value="long">长上下文主实验</option><option value="fresh">七组辅助实验</option><option value="public">真实对话</option><option value="regression">旧问题回归</option></select><label><input type="checkbox" id="issues">只看含问题回答的题</label><span id="count"></span></header>
"""
        + "\n".join(parts)
        + """<script>const items=[...document.querySelectorAll('.case')];function filter(){let n=0;const q=document.querySelector('#search').value.toLowerCase(),s=document.querySelector('#suite').value,b=document.querySelector('#issues').checked;for(const el of items){const show=(!s||el.dataset.suite===s)&&(!q||el.textContent.toLowerCase().includes(q))&&(!b||el.querySelector('.bad'));el.hidden=!show;if(show)n++;}document.querySelector('#count').textContent=`显示 ${n} / ${items.length} 题`;}document.querySelectorAll('input,select').forEach(el=>el.addEventListener('input',filter));filter();</script></html>"""
    )
    path.write_text(page, encoding="utf-8")


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    for name, size in [
        ("Normal", 11),
        ("Title", 25),
        ("Subtitle", 12),
        ("Heading 1", 18),
        ("Heading 2", 12),
    ]:
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.italic = False
        for border in s.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
        s.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.paragraph_format.space_after = Pt(7)
        s.paragraph_format.line_spacing = 1.2
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("实验报告  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)
    return doc


def table(doc, headers, data, widths=None):
    if widths:
        widths = [w * 17 / sum(widths) for w in widths]
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for c, w in zip(t.columns, widths):
            c.width = Cm(w)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = str(text)
    for record in data:
        for cell, text in zip(t.add_row().cells, record):
            cell.text = str(text)
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "233C55" if i == 0 else "F0F4F7" if i % 2 else "FFFFFF")
            pr.append(shade)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                v = OxmlElement("w:" + side)
                v.set(qn("w:w"), "85")
                v.set(qn("w:type"), "dxa")
                margins.append(v)
            pr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                v = OxmlElement("w:" + side)
                v.set(qn("w:val"), "single")
                v.set(qn("w:sz"), "4")
                v.set(qn("w:color"), "D9D9D9")
                borders.append(v)
            pr.append(borders)
            for p in cell.paragraphs:
                if cell.text.replace(".", "", 1).isdigit():
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def chapter(doc, title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def paragraph(doc, text):
    doc.add_paragraph(text)


def build(args):
    args.out.mkdir(parents=True, exist_ok=args.refresh)
    spec = [
        ("long", "long_context_luna_01", "long_context_02"),
        ("fresh", "episode_final_luna_01", "episode_final_01"),
        ("public", "episode_public_eval_01", "episode_public_eval_01"),
        ("regression", "episode_regression_luna_01", "episode_regression_01"),
    ]
    records = [r for s, j, o in spec for r in joined(s, RUNS / j, RUNS / o)]
    assert len(records) == 376
    data = {s: [r for r in records if r["suite"] == s] for s, _, _ in spec}
    summaries = {s: summary(rs) for s, rs in data.items()}
    long_pair = paired(data["long"], "base_full_history")
    metrics = dict(
        suites=summaries,
        long_paired=long_pair,
        long_by_budget={
            str(b): summary([r for r in data["long"] if r["budget"] == b])
            for b in (2048, 8192, 16384, 24576, 32000)
        },
        long_by_seed={
            str(seed): summary([r for r in data["long"] if r["seed"] == seed])
            for seed in (57511, 57512)
        },
        long_by_probe={
            kind: summary([r for r in data["long"] if r["output_id"].endswith("-" + kind)])
            for kind in (
                "early",
                "middle",
                "recent",
                "correction",
                "conflict",
                "missing",
                "stable",
                "absent",
            )
        },
    )
    (args.out / "results.json").write_text(
        json.dumps(dict(metrics=metrics, records=records), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    appendix(records, args.out / "逐题对照与证据.html")
    doc = setup_doc()
    doc.add_paragraph("长上下文下原始模型与磁盘记忆模型", style="Title")
    doc.add_paragraph("同一聊天历史下的对照实验报告", style="Subtitle")
    paragraph(
        doc,
        "实验对象  Qwen2.5 1.5B Instruct\n实验设备  RTX 4060 Laptop 8 GB\n结果整理  2026年10月26日\n主实验  两段独立对话  五个长度检查点  160条回答",
    )
    a = summaries["long"]["base_full_history"]
    b = summaries["long"]["episode_open_grouped"]
    paragraph(
        doc,
        f"本报告直接比较原始模型读取完整长聊天，与同一冻结模型加载可训练磁盘记忆模块后的表现。最长输入约3.2万token，所有原始模型输入均在32768位置上限内，没有截断历史。主实验中，整句合格回答由原始模型的{a['whole'].get('acceptable', 0)}/80变为记忆版的{b['whole'].get('acceptable', 0)}/80；这包含证据不足时的合理不作答，不等于事实全部答对。",
    )
    paragraph(
        doc,
        "结论必须同时看两面：记忆模块可以缩小读取范围，帮助保留早期事实与更正；但它仍可能遗漏事件片段、将聊天对象误当行为人，或在可回答时表达不确定。辅助真实对话测试中，直接读取原文的原始模型仍比当前连续编码更可靠。本轮没有证明通用长期记忆或无限上下文。",
    )
    table(
        doc,
        ["导出内容", "用途"],
        [
            ("本报告", "设计、训练、同题结果、失败原因与结论"),
            ("逐题对照与证据.html", "搜索全部376条回答，查看原文证据及复核理由"),
            ("results.json", "结构化逐条记录与汇总，可独立重算"),
            ("evidence 目录", "原始输出、协议、评审记录、权重与代码快照"),
        ],
        [5, 12],
    )
    paragraph(
        doc,
        "主实验以长上下文对照为中心。前一轮七组实验、真实开源对话测试和旧问题回归仅作辅助证据，不混成一个总准确率。所有正确性结论来自AI逐条证据复核，尚无人工金标准。",
    )

    chapter(doc, "一 公平比较的具体定义")
    paragraph(
        doc,
        "两种系统接收同一条连续增长的用户消息流。每个检查点使用相同问题、相同已发生事件、相同系统行为指令和相同100个新token生成上限，均采用确定性解码。没有在评估期间更新参数，也没有把检查题及其答案写回后续对话或磁盘记忆。",
    )
    table(
        doc,
        ["项目", "原始模型", "加载记忆模型"],
        [
            ("基座权重", "官方1.5B指令模型，冻结", "同一份基座权重，冻结"),
            ("历史来源", "当前检查点的完整聊天原文", "每个用户轮次先编码并落盘"),
            ("回答时读取", "全部已有历史，未截断", "自动匹配事件的连续编码"),
            ("系统指令及问题", "相同", "相同"),
            ("原文访问", "允许，作为完整历史基线", "记忆回答不读取聊天原文"),
            ("保存状态", "全部因果KV，用于执行加速", "SQLite浮点编码及地址索引"),
        ],
        [3.1, 6.95, 6.95],
    )
    paragraph(
        doc,
        "这是系统级比较：可用历史事实相同，回答时的表示和读取范围不同。不能把它解释成两个模型在完全相同token序列上的比较；过滤上下文和训练编码共同构成了修改版的能力。原始模型的KV缓存只是标准计算状态，并没有增加训练后的外部事件记忆。",
    )
    paragraph(
        doc,
        "为了适配本机显存，原始模型增量预填充并保留全部历史KV，每个问题复制独立缓存。它不删除远处token，不采用滑动窗口。2道短例以及3道实际长输入（8K一例、32K两例）的缓存执行与一次性完整预填充回答均逐字一致，且3项因果注意力检查通过；这不是所有问题的数值等价证明。耗时须分别报告历史预处理与问答时间，不能省略预处理后宣称速度优势。",
    )

    chapter(doc, "二 事件记忆模块怎样工作")
    paragraph(
        doc,
        "写入流程为 用户片段 → 冻结基座输入嵌入 → 可训练宽度投影及局部卷积 → float16连续编码 → SQLite。读取流程为 问题地址编码 → 动态簇候选 → 地址一致性检查 → 同事件片段按顺序聚合 → 嵌入解码 → 冻结语言模型生成。",
    )
    table(
        doc,
        ["组件", "实际实现与边界"],
        [
            ("连续编码", "1536维输入投影到256维，深度逐通道卷积保留局部上下文，解码回1536维"),
            ("可训练规模", "编码与解码共787200参数，基座参数不更新"),
            ("事件归组", "动态在线簇辅助候选选择，必要时扫描索引；同地址片段按写入先后读取"),
            ("更新语义", "不分类为补充或更正后覆盖；旧说法和新说法都保留，由语言模型解释"),
            ("持久化", "SQLite保存float16向量、地址、来源与顺序；运行时不保存原文或token ID"),
            (
                "容量限制",
                "每片段最多192 token，一次事件读取最多4096 token；地址仍为16项目乘16讨论次数",
            ),
        ],
        [3.2, 13.8],
    )
    paragraph(
        doc,
        "这版没有六字段答案分类门控，也没有固定拒答模板。地址匹配门槛仍为0.99，但旧版字段支持概率门槛已不属于新路径，因此不能声称新旧具有相同的拒答保证。",
    )
    paragraph(
        doc,
        "它压缩的是嵌入宽度，而不是事件时长。每token编码约512字节，通常比原始文字更大，也可能恢复原措辞。随着记忆增长，磁盘内容和地址索引仍增长。因此本轮验证的是可学习、可落盘的信息通道，尚不是人类式抽象记忆。",
    )

    chapter(doc, "三 训练数据和损失函数")
    paragraph(
        doc,
        "语言能力来自已训练好的Qwen基座。公开语料训练主要改善记忆通道保真，没有重新训练一个会说话的1.5B模型。候选材料包括FineWeb2中文网页与OpenAssistant OASST2对话。按整篇文档URL或整棵对话树划分训练、验证、测试，只有标准化文本精确去重，不保证消除全部近重复。",
    )
    table(
        doc,
        ["数据口径", "数量"],
        [
            ("清洗后正文", "3221541544字节，853248条"),
            ("训练池", "819055条，约3.09 GB"),
            ("验证与测试", "17243条验证，16950条测试"),
            ("抽样池", "24000篇网页加8000条对话，5892570 token"),
            ("实际不同训练片段", "26768个，共4930827 token"),
            ("公开材料累计暴露", "10597138 token，包含重复抽样"),
        ],
        [6, 11],
    )
    paragraph(
        doc,
        "初始控制版使用512组合成经历、896片段，覆盖单次叙述、更正、质疑、补充。先进行400步嵌入重建，再做120步已知事实问答与重建联合优化；384个正向可答案例可被抽取，不确定案例没有固定拒答句监督。",
    )
    paragraph(
        doc,
        "重建损失 Lrec 是重建嵌入与原始嵌入的均方误差，除以输入嵌入的平均平方幅度。初始问答阶段使用 Lqa 加 2倍Lrec。公开材料阶段做2400步重建，每步24个公开片段加8个合成片段；AdamW学习率0.0001。之后80步以冻结基座的下一token分布为教师，损失为 KL教师到学生 加 2倍Lrec，学习率0.00001，每段监督12个位置。",
    )
    paragraph(
        doc,
        "验证相对嵌入MSE从0.6511降到0.4857，训练峰值显存约3.79 GB。这个指标反映信息重建而非回答正确性。没有把全部3.22 GB用于梯度更新，也没有训练固定拒答句。",
    )

    chapter(doc, "四 长上下文主实验的数据构造")
    paragraph(
        doc,
        "采用种子57511与57512生成两段独立对话。关键经历包含明确人物、事件、感受、原因和关注点，使用新的组合与地址；中间穿插来自保留测试集的中文阅读段落。用户消息与固定助手确认交替出现，每段对话最终213个用户轮次、426条消息。阅读段落只作为干扰材料，不被标注为用户亲历事件。",
    )
    lengths = []
    for seed in (57511, 57512):
        for cp in read(RUNS / "long_context_02" / f"seed_{seed}" / "protocol.json")[
            "token_checkpoints"
        ]:
            lengths.append(
                [
                    seed,
                    cp["budget"],
                    f"{cp['question_input_token_range'][0]}至{cp['question_input_token_range'][1]}",
                    cp["user_turns"],
                    cp["event_fragments"],
                    cp["reading_fragments"],
                ]
            )
    table(
        doc,
        ["种子", "名义长度", "实际输入token", "用户轮次", "事件片段", "阅读片段"],
        lengths,
        [2.1, 2.3, 4.3, 2.7, 2.8, 2.8],
    )
    paragraph(
        doc,
        "每个检查点有8类问题：最早事件人物、中段事件内容、最近事件关注点、原因更正、未核实情绪分歧、已知事件缺少人物、稳定情绪对照、完全未出现的事件。更正在约6K位置写入，分歧在约12K位置写入，因此能分别检查更新前和更新后的行为。",
    )
    paragraph(
        doc,
        "这是受控事件加真实阅读材料的长聊天压力测试，不是完整自然人类聊天，也没有开放域事件边界识别。两段对话和跨检查点重复题不能视为80个独立随机样本。本轮不使用显著性检验或宣称总体泛化准确率。",
    )

    chapter(doc, "五 随上下文增长的主要结果")
    table(
        doc,
        ["长度", "原始整句合格", "记忆整句合格", "原始有问题", "记忆有问题"],
        [
            [
                b,
                metrics["long_by_budget"][str(b)]["base_full_history"]["whole"].get(
                    "acceptable", 0
                ),
                metrics["long_by_budget"][str(b)]["episode_open_grouped"]["whole"].get(
                    "acceptable", 0
                ),
                metrics["long_by_budget"][str(b)]["base_full_history"]["whole"].get(
                    "problematic", 0
                ),
                metrics["long_by_budget"][str(b)]["episode_open_grouped"]["whole"].get(
                    "problematic", 0
                ),
            ]
            for b in (2048, 8192, 16384, 24576, 32000)
        ],
        [2, 3.8, 3.8, 3.7, 3.7],
    )
    paragraph(
        doc,
        "每个长度每种模型16条回答，来自两段对话各8题。整句合格包括回答正确以及确实缺少证据时合理保留不确定；有证据却没有回答应另记为未完成。目标正确但附加错误行为人或因果，不算整句合格。",
    )
    table(
        doc,
        ["总计各80条", "目标正确", "整句合格", "有问题", "诚实未完成", "事实问题"],
        [
            [
                MODES[m],
                s["target"].get("correct", 0),
                s["whole"].get("acceptable", 0),
                s["whole"].get("problematic", 0),
                s["whole"].get("honest_but_incomplete", 0),
                s["event_errors"],
            ]
            for m, s in summaries["long"].items()
        ],
        [4.5, 2.3, 2.5, 2.2, 3, 2.5],
    )
    paragraph(
        doc,
        f"按同一道题配对，双方均合格{long_pair['both_good']}题，仅记忆版合格{long_pair['new_only']}题，仅原始模型合格{long_pair['base_only']}题，双方均未合格{long_pair['both_bad']}题。比较的是整句可接受性，未合格不全等于编造；完整标签和原始回答见HTML附录。",
    )
    paragraph(
        doc,
        "长度增加同时也增加了干扰材料、事件数量及事实距离；因此这组曲线衡量的是增长对话下的系统表现，不能单独归因为位置长度。若要严格分离因果，应固定事件集合，仅改变间隔与填充内容，并做多随机种子与多种位置布置。",
    )
    paragraph(
        doc,
        "两段对话的方向一致，但幅度不同：种子57511的整句合格数从13/40升到36/40，种子57512从20/40升到30/40。原始模型并非完全遗忘早期信息，最早人物题仍有9/10合格；更明显的差异出现在中段事件、稳定情绪和缺失字段处理。",
    )

    chapter(doc, "六 逐题对照与失败机制")
    examples = [
        (57511, "L16384-correction"),
        (57511, "L32000-middle"),
        (57511, "L08192-conflict"),
    ]
    for seed, oid in examples:
        rs = [r for r in data["long"] if r["seed"] == seed and r["output_id"] == oid]
        if not rs:
            continue
        doc.add_heading(f"对话 {seed} 题目 {oid}", 2)
        paragraph(doc, rs[0]["packet"]["question"])
        for r in sorted(rs, key=lambda x: x["mode"]):
            paragraph(
                doc,
                f"{MODES[r['mode']]}：{r['packet']['answer']}\n复核：{LABELS[r['review']['whole_reply']]}。{r['review']['rationale']}",
            )
    paragraph(
        doc,
        "这些例子分别观察更正保留、额外主体错误及不必要的不确定。不能只匹配答案里的关键词；‘提前完成’与‘小何提前完成任务’对于没有明确行为人的证据并不等价。",
    )

    chapter(doc, "七 读取证据与答对结果要分开")
    paragraph(
        doc,
        "报告同时保存三种信息：该检查点本应可用的相关事件原文，记忆路径实际读取的片段ID，以及最后交付的回答。原文只供基线和评审使用，不是注入记忆模型的答案。实际读取不全时偶然答对，不能证明系统掌握了全部更新。",
    )
    paragraph(
        doc,
        "主实验记忆版80次读取中，77次片段ID集合与应有证据一致（包括没有匹配事件时读取空集合）。另外3次遗漏T0081分歧片段，却仍回答无法确定；这些回答在输出层合格，但不能算成功读懂冲突。其余14次未合格回答都已读取对应证据：5次擅自指定行为人或增加因果，5次将缺失的交流对象姓名回答为‘用户’，4次在分歧尚未出现、心情明确时回答无法确定。完整读取也不保证理解正确。",
    )
    paragraph(
        doc,
        "失败可分为地址或检索未覆盖、连续编码失真、语言模型对关系理解错误、输出遗漏，以及对证据充分性的误述。仅有完整事件原文不能证明模型实际读到了它；反过来，读取了两段冲突记忆也不保证模型会保留分歧。",
    )
    paragraph(
        doc,
        "程序已移除旧答案模板门控。因此新模型自由生成错误，应归入新读取或生成路径，不能再归咎于旧模板。仍保留地址归组范围和0.99匹配门槛；实验没有为获得更多回答而调低这一阈值。",
    )
    runtimes = []
    for seed in (57511, 57512):
        for cp_index, cp in enumerate(
            read(RUNS / "long_context_02" / f"seed_{seed}" / "checkpoint_runtime.json")
        ):
            if cp_index not in (0, 2, 4):
                continue
            runtimes.append(
                [
                    seed,
                    cp["input_tokens"],
                    f"{cp['prefill_seconds']:.2f}",
                    f"{cp['writing_seconds']:.2f}",
                    f"{cp['peak_allocated_gb']:.2f}",
                    f"{cp['store']['file_bytes'] / 1e6:.2f}",
                ]
            )
    table(
        doc,
        ["种子", "历史token", "增量预填秒", "增量写入秒", "合并峰值GB", "数据库MB"],
        runtimes,
        [2, 2.8, 3.3, 3.3, 3, 2.6],
    )
    paragraph(
        doc,
        "峰值是在两种系统同一进程内运行的合并值，不是单独模型显存排名。预填充与写入时间只计该检查点新增历史，不能视为重新处理全部历史的耗时。最终SQLite是最后检查点状态，没有保存每个检查点的独立数据库快照。",
    )

    chapter(doc, "八 前一轮七组对比作为辅助分析")
    paragraph(
        doc,
        "前一轮最长约8K输入，使用另一个新事件种子和改写问题，覆盖8、32、64、128个用户轮次。它用于区分旧门控、数据扩展与归组修改，不能替代本报告的长上下文主实验。",
    )
    table(
        doc,
        ["条件 各24条", "目标正确", "整句合格", "有问题", "未完成"],
        [
            [
                MODES[m],
                s["target"].get("correct", 0),
                s["whole"].get("acceptable", 0),
                s["whole"].get("problematic", 0),
                s["whole"].get("honest_but_incomplete", 0),
            ]
            for m, s in summaries["fresh"].items()
        ],
        [7, 2.5, 2.5, 2.5, 2.5],
    )
    paragraph(
        doc,
        "公开语料版相较小数据版，三个明确更正后问题由旧原因转为新原因，但更正前的一题变成答非所问，另有‘解决故障’读成‘了解故障’。新增地址训练没有提高最终合格数量，新表述下的分歧片段仍可能遗漏。",
    )
    p = paired(data["fresh"], "base_same_event_text")
    paragraph(
        doc,
        f"最接近读取通道的辅助控制是原始基座读取同一批实际检索片段原文：双方均合格{p['both_good']}题，仅记忆版合格{p['new_only']}题，仅原文组合格{p['base_only']}题，双方均未合格{p['both_bad']}题。片段语义内容一致，但输入嵌入及包装token并不完全一致；它不是随机重复实验，也不能证明编码普遍优于原文。",
    )

    chapter(doc, "九 真实开源对话给出的反面证据")
    paragraph(
        doc,
        "从OASST2保留测试集选取8棵不同对话树，中文和英文各4。每例2至4个片段，要求概括用户明确说过的需求和限制，不能把助手建议当作用户决定。该测试给定对话归属，只检验编码与读取，未测试自动开放域归组。",
    )
    table(
        doc,
        ["条件 各8条", "完整正确", "部分对应", "目标无依据", "事实问题"],
        [
            [
                MODES[m],
                s["target"].get("correct", 0),
                s["target"].get("partial", 0),
                s["target"].get("unsupported_answer", 0),
                s["event_errors"],
            ]
            for m, s in summaries["public"].items()
        ],
        [7, 2.5, 2.5, 2.5, 2.5],
    )
    paragraph(
        doc,
        "原始模型直接读取原文有3条完整正确、5条不完整，但没有新增无依据事实；公开语料记忆版仍无完整正确摘要，5条含事实问题。结果表明记忆通道仍会损伤开放域实体、用途、限制和说话来源。受控长对话的优势不能推广为开放聊天优势。",
    )
    for cid in ("P00", "P07"):
        rs = [r for r in data["public"] if r["case_id"] == cid]
        doc.add_heading(f"真实对话案例 {cid}", 2)
        for r in sorted(rs, key=lambda x: x["mode"]):
            paragraph(doc, f"{MODES[r['mode']]}：{r['packet']['answer']}")
    paragraph(
        doc,
        "这里的摘要完整性有主观判断空间，故保留部分对应、不完整和事实错误三种信息，避免将不够详细与捏造混成一个分数。",
    )

    chapter(doc, "十 旧问题回归和普通聊天")
    paragraph(
        doc,
        "在原来已知的24道连续聊天问题上，公开语料及归组候选得到18次目标正确、4次因人物未记录而合理不作答、2次擅自选择未核实情绪。原先5次正确内容被模板拦截和3次明确原因更正未采用，均在这些已知用例上得到正确输出。",
    )
    paragraph(
        doc,
        "仍失败的两条都实际读取了高兴和生气两版记忆，却直接认定生气。这是理解与生成问题，说明保留所有片段只是必要条件，不足以保证正确处理不确定性。回归成绩不能当作未见泛化成绩。",
    )
    paragraph(
        doc,
        "普通聊天通过显式入口绕过记忆，仍使用冻结基座。4个普通聊天测试在挂载模块前后输出逐字一致，覆盖介绍、解释下雨、翻译和乘法。它证明这条旁路没有改变输出，不代表已实现自动记忆调用，也不代表广泛聊天质量评测。",
    )
    paragraph(
        doc,
        "相关工程检查覆盖磁盘重启、顺序、说话来源、所有者隔离、事务回滚、旧数据库兼容，以及引用校验。此前32项单元检查通过；本轮额外验证分块注意力保留完整因果上下文。代码格式检查与语义正确性评审是不同的验证层次。",
    )

    chapter(doc, "十一 评审流程和可审计性")
    paragraph(
        doc,
        "GPT 5.6 Luna按条读取问题、回答、该时点相关事件证据及执行日志进行初审，组别和长度隐藏。主助手在读取Luna判断前独立核对证据，随后校验引用是否确实存在、哈希是否绑定原始回答、是否覆盖所有条目、断言判断与总标签是否冲突。所有分歧复核，并随机抽查一致项。",
    )
    audit_rows = []
    for suite, j, _ in spec:
        au = read(RUNS / j / "audit.json")
        audit_rows.append(
            [
                suite,
                au["reviewed"],
                au["disagreements"],
                au["quality_flags"],
                len(au["agreement_sample"]),
            ]
        )
    table(
        doc, ["实验", "逐条数", "分歧数", "一致性标志", "一致项抽查"], audit_rows, [5, 3, 3, 3, 3]
    )
    paragraph(
        doc,
        "标签区分目标完成、整句可靠性、事件事实问题和状态问题。自然表达的不确定不需要符合固定句式；有证据却拒答不算答对；正确姓名后面附加错误主体仍可能整句不合格。是非问题按其实际语义判断，不能机械把‘不能确定’都标成未回答。",
    )
    paragraph(
        doc,
        "所有标签均为AI判断与AI复核，不是人工金标准。主助手参与了实验开发，并非完全盲评；独立Luna评审也可能从回答样式推断条件。评审一致不等于真实正确，随机抽查不能排除未发现错误。导出原始判断与复核记录，便于后续人工审查。",
    )

    chapter(doc, "十二 结论和下一步研究边界")
    paragraph(
        doc,
        "在这两段受控长对话中，加载记忆模块有明确的系统级改善：整句合格由33/80升至66/80，有问题回答由39/80降至10/80，未完成由8/80降至4/80。约32K处由5/16升至13/16。同题有35次改善、2次退步；这里的改善指整句合格性，不代表所有题都提供了具体事实。系统读取范围缩小可以降低干扰，但仍引入归组遗漏与编码失真；它没有让基座本身拥有无限上下文。",
    )
    paragraph(
        doc,
        "当前最重要的约束包括：只有两段受控长对话；地址空间封闭；输入片段按用户轮次切分；记忆保留序列长度而非抽象概念；读取有4096 token预算；索引仍随记录增长；真实开放对话保真不足；没有人体长期记忆或情感意识层面的验证。",
    )
    paragraph(
        doc,
        "后续应分别检验三条路径。第一，固定事件集合而改变间隔和位置，以区分长度、干扰与检索难度。第二，用真实连续对话训练实体、否定、时间、说话来源及更正关系的保真目标，并允许自然表达不确定。第三，学习开放域事件关联与可验证的摘要合并，同时保留历史修订来源。每次改变都保留原始完整上下文基线和同片段原文控制。",
    )
    paragraph(
        doc,
        "在证据没有支持以前，不把目标定为无限记忆，也不以调低拒答门槛换取更多输出。优先减少具体事实错误、错误归因和漏读修订，再讨论更高压缩率或更大规模。",
    )

    chapter(doc, "十三 文件索引与复现说明")
    table(
        doc,
        ["路径", "说明"],
        [
            ("results.json", "376条回答、证据、评审及派生汇总"),
            ("逐题对照与证据.html", "可搜索、按实验筛选的离线逐题对照"),
            ("evidence/long_context_02", "两段原始对话、问题、输出、窗口长度及运行日志"),
            ("evidence/long_context_luna_01", "盲审包、Luna原判、主审、复核、随机抽查"),
            ("evidence/episode_memory_open_01", "公开语料训练参数、日志、采样来源及编码权重"),
            ("evidence/code_snapshot", "导出时的源码快照及环境要求"),
            ("MANIFEST.json", "导出文件大小和SHA256校验"),
        ],
        [7.1, 9.9],
    )
    paragraph(
        doc,
        "复现时需要原始Qwen2.5-1.5B-Instruct权重及Python依赖；完整基座和数GB原始语料不放入结果包。源数据版本、许可和下载校验保存在协议中。训练材料来源为https://huggingface.co/datasets/HuggingFaceFW/fineweb-2与https://huggingface.co/datasets/OpenAssistant/oasst2。",
    )
    paragraph(
        doc,
        "主实验入口为 experiments.episodic_memory.long_context_comparison；用新输出目录运行，不覆盖既有结果。模型目录使用episode_memory_grouped_01，语料指向corpora/open_01/prepared。正式实验目录long_context_02；long_context_01只发生方法初始化错误，没有正式测试输出，不混入统计。",
    )
    paragraph(
        doc,
        "缓存敏感性检查见long_cache_sensitivity.json。评估与导出过程不重训权重。源码快照包含导出时实现，历史协议及实际输出是本轮实验事实的依据。原始记录若有局部表述修正，保留审计说明而不覆盖旧评审。",
    )
    doc.save(args.out / "长上下文记忆模型实验报告.docx")
    evidence = args.out / "evidence"
    for _, j, o in spec:
        for name in set((j, o)):
            dest = evidence / name
            dest.mkdir(parents=True, exist_ok=True)
            for p in (RUNS / name).rglob("*"):
                if p.is_file() and p.suffix in (".json", ".jsonl", ".txt", ".sqlite", ".py"):
                    target = dest / p.relative_to(RUNS / name)
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(p, target)
    for name in (
        "episode_memory_02",
        "episode_memory_open_01",
        "episode_memory_grouped_01",
        "episode_identity_01",
        "fact_reader_01",
    ):
        dest = evidence / name
        dest.mkdir(parents=True, exist_ok=True)
        for filename in (
            "metadata.json",
            "training_log.json",
            "training_sample.json",
            "validation_sample.json",
            "codec.pt",
            "address.pt",
            "adapter.pt",
            "semantic.pt",
            "train.json",
            "validation.json",
        ):
            p = RUNS / name / filename
            if p.exists():
                shutil.copy2(p, dest / filename)
    code = evidence / "code_snapshot"
    code.mkdir(parents=True, exist_ok=True)
    for p in ROOT.glob("*.py"):
        shutil.copy2(p, code / p.name)
    for p in ROOT.glob("requirements*.txt"):
        shutil.copy2(p, code / p.name)
    for p in ("EPISODE_MEMORY_PROTOCOL.md", "EPISODE_MEMORY_RESULTS.md"):
        shutil.copy2(ROOT / p, evidence / p)
    shutil.copy2(
        RUNS / "episode_chat_invariance_01.json", evidence / "episode_chat_invariance_01.json"
    )
    shutil.copy2(
        ROOT / "corpora/open_01/prepared/manifest.json", evidence / "public_corpus_manifest.json"
    )
    (args.out / "阅读说明.txt").write_text(
        "先阅读PDF或Word正式报告，再用HTML逐题检查。results.json是结构化结果，evidence保留证据与实验配置。\n主实验160条，辅助七组168条，真实对话24条，旧回归24条，共376条。\n所有正确性标签为AI评审，非人工金标准。权重只包含小型记忆模块，不含完整Qwen基座或数GB语料。\n",
        encoding="utf-8",
    )
    print(
        json.dumps(
            dict(
                output=str(args.out), records=len(records), main_metrics=metrics["long_by_budget"]
            ),
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--refresh", action="store_true")
    build(p.parse_args())
